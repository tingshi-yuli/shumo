"""Portable result reproduction from the frozen canonical trajectories.
Only four-decimal exports; no invented version-comparison tables.
"""
import gzip
import json
import shutil
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from model import sample

ROOT=Path(__file__).resolve().parent

def write_book(name,sheets):
    wb=Workbook(write_only=True)
    for title,states,field,moving in sheets:
        ws=wb.create_sheet(title)
        ws.freeze_panes='B2';ws.column_dimensions['A'].width=25
        for j in range(2,25):ws.column_dimensions[get_column_letter(j)].width=12
        header=['时间/s；距离/cm']+[round(x/10,1) for x in range(21)]
        if moving:header+=['药材表面','表面半径/cm']
        ws.append(header)
        for state in states:
            vals=sample(state,field,np.arange(21)/10)
            if moving:vals+=[state[field][-1],state['radius_cm']]
            values=[state['t']]+[None if v is None else round(v,4) for v in vals]
            cells=[WriteOnlyCell(ws,value=v) for v in values]
            for c in cells:c.number_format='0.0000'
            ws.append(cells)
    path=ROOT/name
    wb.save(path)
    (ROOT/'results').mkdir(exist_ok=True)
    shutil.copy2(path,ROOT/'results'/name)

def main():
    for name,number in [('q1_final',1),('q2_full',2),('q4_final',4)]:
        with gzip.open(ROOT/'runs'/f'{name}.json.gz','rt',encoding='utf-8') as f:
            result=json.load(f)
        if number in (1,2):
            write_book(f'result{number}.xlsx',[(title,result['short'],field,False)
                for title,field in [('温度','T'),('水分浓度','C')]])
            if number==2:write_book('result3.xlsx',[('Sheet1',result['long'],'C',False)])
        else:write_book('result4.xlsx',[('Sheet1',result['long'],'C',True)])
        print('exported',name,flush=True)
        del result

if __name__=='__main__':main()
